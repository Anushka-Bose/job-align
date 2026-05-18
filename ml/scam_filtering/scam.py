import hashlib
import json
import os
import re
from datetime import date
from typing import Dict, List

from dotenv import load_dotenv

load_dotenv()

try:
    from google import genai
except Exception:
    genai = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    import fitz
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None


SECTION_HEADERS = [
    "experience", "work history", "employment",
    "education", "skills", "projects",
    "certifications", "summary", "objective"
]

LLM_MODEL = os.environ.get("SCAM_FILTER_MODEL", "gemini-3.1-flash-lite")
MAX_LLM_CHARS = 2600
_client = None
_llm_unavailable = False
_llm_cache: Dict[str, Dict] = {}
_last_llm_error = ""
PROMPT_VERSION = "chronology_guard_v2"

MONTHS = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}

DATE_RANGE_RE = re.compile(
    r"((?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\s+)?(?:19|20)\d{2})"
    r"\s*(?:-|to|through|\u2013|\u2014)\s*"
    r"((?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\s+)?(?:19|20)\d{2}|present|current|now)",
    re.IGNORECASE
)

CHRONOLOGY_NEGATION_TOKENS = (
    "future relative to the current date",
    "future relative to current date",
    "future-dated",
    "in the future",
    "impossible timeline",
    "timeline is illogical",
    "chronology issue",
    "chronological inconsistency",
    "significant resume error",
)


def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""

    if ext == ".pdf":
        if fitz is not None:
            try:
                with fitz.open(file_path) as pdf:
                    for page in pdf:
                        extracted = page.get_text()
                        if extracted:
                            text += extracted + " "
            except Exception as e:
                return f"ERROR reading PDF: {e}"
        elif PyPDF2 is not None:
            try:
                with open(file_path, "rb") as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages:
                        extracted = page.extract_text()
                        if extracted:
                            text += extracted + " "
            except Exception as e:
                return f"ERROR reading PDF: {e}"
        else:
            return "ERROR: No PDF reader installed"

    elif ext in [".docx", ".doc"]:
        if docx is None:
            return "ERROR: python-docx not installed"
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + " "
        except Exception as e:
            return f"ERROR reading DOCX: {e}"

    elif ext == ".txt":
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        except Exception as e:
            return f"ERROR reading TXT: {e}"
    else:
        return "ERROR: Unsupported file format"

    return text.strip()


def check_essential_info(text):
    email = bool(re.search(r"[\w.%+-]+@[\w.-]+\.[a-zA-Z]{2,}", text))
    phone = bool(re.search(r"\+?\d[\d \-\(\)\.]{8,15}\d", text))
    links = bool(re.search(r"https?://[^\s]+|www\.[^\s]+", text))
    return email, phone, links


def check_structure(text):
    lowered = text.lower()
    return [header for header in SECTION_HEADERS if re.search(r"\b" + re.escape(header) + r"\b", lowered)]


def _risk(score: int) -> str:
    if score >= 70:
        return "HIGH"
    if score >= 40:
        return "MEDIUM"
    return "LOW"


def _clamp_score(value: float) -> int:
    return int(max(0, min(100, round(value))))


def _llm_enabled() -> bool:
    setting = os.environ.get("SCAM_FILTER_USE_LLM", "auto").lower()
    if setting in ("0", "false", "no", "off", "disabled"):
        return False
    return bool(genai is not None and os.environ.get("GEMINI_API_KEY"))


def _is_llm_error(error: Exception) -> bool:
    msg = str(error).lower()
    return any(token in msg for token in (
        "resource_exhausted", "quota exceeded", "429",
        "no connection", "connection", "timeout", "api key",
        "rate limit", "unavailable"
    ))


def _get_client():
    global _client
    if genai is None:
        raise RuntimeError("google-genai is not installed")
    if _client is None:
        api_key = os.environ.get("GEMINI_API_KEY")
        if not api_key:
            raise RuntimeError("GEMINI_API_KEY not configured")
        _client = genai.Client(api_key=api_key)
    return _client


def _strip_json_fence(content: str) -> str:
    text = (content or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\s*```$", "", text)
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end >= start:
        return text[start:end + 1]
    return text


def _parse_month_year(token: str, today: date) -> tuple[int, int] | None:
    text = (token or "").strip().lower()
    if text in ("present", "current", "now"):
        return today.year, today.month

    year_match = re.search(r"\b(19|20)\d{2}\b", text)
    if not year_match:
        return None

    year = int(year_match.group(0))
    month = 1
    for name, number in MONTHS.items():
        if re.search(rf"\b{name}\b", text):
            month = number
            break
    return year, month


def _extract_chronology(text: str, today: date | None = None) -> Dict:
    today = today or date.today()
    entries = []
    lines = [line.strip() for line in re.split(r"[\r\n]+", text or "") if line.strip()]

    for index, line in enumerate(lines):
        for match in DATE_RANGE_RE.finditer(line):
            context = " ".join(lines[max(0, index - 1): min(len(lines), index + 2)]).lower()
            category = "other"
            if re.search(r"\b(education|b\.?tech|bachelor|undergraduate|university|college|degree|cgpa)\b", context):
                category = "education"
            elif re.search(r"\b(intern|internship|trainee)\b", context):
                category = "internship"
            elif re.search(r"\b(experience|engineer|developer|analyst|associate|work)\b", context):
                category = "experience"

            start = _parse_month_year(match.group(1), today)
            end = _parse_month_year(match.group(2), today)
            if not start or not end:
                continue

            entries.append({
                "category": category,
                "start": start,
                "end": end,
                "raw": match.group(0),
                "context": context,
            })

    issues: List[str] = []
    future_dates: List[str] = []
    valid_student_internship = False
    today_value = (today.year, today.month)

    for entry in entries:
        if entry["start"] > entry["end"]:
            issues.append(f"date range is reversed: {entry['raw']}")
        if entry["start"] > today_value:
            future_dates.append(entry["raw"])
        if entry["end"] > today_value and not re.search(r"\b(present|current|now)\b", entry["raw"], re.IGNORECASE):
            future_dates.append(entry["raw"])

    education_entries = [entry for entry in entries if entry["category"] == "education"]
    internship_entries = [entry for entry in entries if entry["category"] == "internship"]

    for education in education_entries:
        for internship in internship_entries:
            if internship["start"] >= education["start"] and internship["end"] <= today_value:
                valid_student_internship = True
                break
        if valid_student_internship:
            break

    has_issue = bool(issues or future_dates)
    if has_issue:
        summary = "; ".join(issues + [f"future-dated entry: {value}" for value in future_dates[:3]])
    elif valid_student_internship:
        summary = "Timeline appears internally consistent. Internship dates fit within an ongoing education timeline and are not future-dated."
    else:
        summary = "No explicit chronology contradiction detected from parsed date ranges."

    return {
        "entries": entries,
        "issues": issues,
        "future_dates": future_dates,
        "valid_student_internship": valid_student_internship,
        "has_issue": has_issue,
        "summary": summary,
    }


def _sanitize_chronology_claims(result: Dict, chronology: Dict) -> Dict:
    if chronology.get("has_issue"):
        return result

    cleaned = dict(result)
    indicators = [
        indicator
        for indicator in list(cleaned.get("scam_indicators") or [])
        if not any(token in indicator.lower() for token in CHRONOLOGY_NEGATION_TOKENS)
    ]
    cleaned["scam_indicators"] = indicators

    reasoning = str(cleaned.get("scam_reasoning") or cleaned.get("scam_explanation") or "")
    lower_reasoning = reasoning.lower()
    if any(token in lower_reasoning for token in CHRONOLOGY_NEGATION_TOKENS):
        cleaned["scam_reasoning"] = chronology["summary"]
        cleaned["scam_explanation"] = chronology["summary"]
        cleaned["scam_percentage"] = min(int(cleaned.get("scam_percentage", 0) or 0), 25)
        cleaned["scam_risk"] = _risk(cleaned["scam_percentage"])
        cleaned["scam_risk_level"] = cleaned["scam_risk"]

    return cleaned


def _repeat_bullet_ratio(lines: List[str]) -> float:
    starts = []
    for line in lines:
        words = re.findall(r"\b[a-zA-Z]+\b", line.lower())
        if words:
            starts.append(words[0])
    if len(starts) < 4:
        return 0.0
    return max(starts.count(word) for word in set(starts)) / len(starts)


def rule_based_analysis(text: str, method: str = "rules") -> Dict:
    lowered = text.lower()
    words = re.findall(r"\b\w+\b", lowered)
    total_words = len(words)
    unique_ratio = len(set(words)) / total_words if total_words else 0
    lines = [line.strip() for line in re.split(r"[\n\r]+|(?<=\.)\s+", text) if line.strip()]
    email, phone, links = check_essential_info(text)
    sections = check_structure(text)

    scam_indicators = []
    scam_score = 5
    if not email:
        scam_indicators.append("missing email")
        scam_score += 15
    if not phone:
        scam_indicators.append("missing phone")
        scam_score += 15
    if len(sections) < 3:
        scam_indicators.append("few standard resume sections")
        scam_score += 12
    if total_words < 120:
        scam_indicators.append("very short resume")
        scam_score += 10
    if not links:
        scam_indicators.append("no profile/project links")
        scam_score += 6

    suspicious_phrases = [
        "guaranteed job", "pay to apply", "processing fee",
        "no interview", "instant hiring", "registration fee",
        "fake experience", "proxy interview", "will pay after selection"
    ]
    for phrase in suspicious_phrases:
        if phrase in lowered:
            scam_indicators.append(f"suspicious phrase: {phrase}")
            scam_score += 20

    ai_indicators = []
    ai_score = 10
    ai_markers = [
        "as an ai language model", "results-driven", "proven track record",
        "dynamic professional", "leveraging cutting-edge", "passionate and dedicated",
        "highly motivated individual", "seamlessly collaborate"
    ]
    for marker in ai_markers:
        if marker in lowered:
            ai_indicators.append(f"generic/AI-like phrase: {marker}")
            ai_score += 12
    if total_words > 80 and unique_ratio < 0.35:
        ai_indicators.append("low vocabulary variety")
        ai_score += 18
    if _repeat_bullet_ratio(lines) >= 0.45:
        ai_indicators.append("many bullets begin with the same wording")
        ai_score += 12
    if re.search(r"\[[^\]]+\]|\{[^}]+\}", text):
        ai_indicators.append("template placeholders detected")
        ai_score += 20

    scam_score = _clamp_score(scam_score)
    ai_score = _clamp_score(ai_score)

    scam = {
        "score": scam_score,
        "risk_level": _risk(scam_score),
        "indicators": scam_indicators,
        "reasoning": "Checked contact info, sections, length, links, and suspicious hiring/fraud phrases.",
        "method": method
    }
    ai = {
        "score": ai_score,
        "percentage": ai_score,
        "risk_level": _risk(ai_score),
        "indicators": ai_indicators,
        "reasoning": "Estimated from generic phrasing, repetition, vocabulary variety, and template markers.",
        "method": method
    }
    return _compact_result(ai, scam, llm_used=False, llm_status=method)


def _compact_result(ai: Dict, scam: Dict, llm_used: bool, llm_status: str) -> Dict:
    ai_score = _clamp_score(ai.get("percentage", ai.get("score", 0)))
    scam_score = _clamp_score(scam.get("score", 0))
    return {
        "ai_generated_percentage": ai_score,
        "ai_risk": str(ai.get("risk_level") or _risk(ai_score)).upper(),
        "ai_risk_level": str(ai.get("risk_level") or _risk(ai_score)).upper(),
        "ai_indicators": list(ai.get("indicators") or [])[:8],
        "scam_percentage": scam_score,
        "scam_risk": str(scam.get("risk_level") or _risk(scam_score)).upper(),
        "scam_risk_level": str(scam.get("risk_level") or _risk(scam_score)).upper(),
        "scam_indicators": list(scam.get("indicators") or [])[:8],
        "ai_explanation": ai.get("reasoning", ""),
        "ai_reasoning": ai.get("reasoning", ""),
        "scam_explanation": scam.get("reasoning", ""),
        "scam_reasoning": scam.get("reasoning", ""),
        "llm_used": llm_used,
        "llm_status": llm_status
    }


def _normalize_llm_result(parsed: Dict, fallback: Dict) -> Dict:
    ai_raw = parsed.get("ai_generated", {})
    scam_raw = parsed.get("scam_applicant", {})

    if not ai_raw:
        ai_raw = {
            "score": parsed.get("ai_generated_percentage", fallback["ai_generated_percentage"]),
            "risk_level": parsed.get("ai_risk", parsed.get("ai_risk_level", "LOW")),
            "indicators": parsed.get("ai_indicators", []),
            "reasoning": parsed.get("ai_explanation", parsed.get("ai_reasoning", fallback["ai_explanation"]))
        }
    if not scam_raw:
        scam_raw = {
            "score": parsed.get("scam_percentage", parsed.get("scam_score", fallback["scam_percentage"])),
            "risk_level": parsed.get("scam_risk", parsed.get("scam_risk_level", "LOW")),
            "indicators": parsed.get("scam_indicators", []),
            "reasoning": parsed.get("scam_explanation", parsed.get("scam_reasoning", fallback["scam_explanation"]))
        }

    ai_score = _clamp_score(ai_raw.get("percentage", ai_raw.get("score", fallback["ai_generated_percentage"])))
    scam_score = _clamp_score(scam_raw.get("score", fallback["scam_percentage"]))

    ai = {
        "score": ai_score,
        "percentage": ai_score,
        "risk_level": str(ai_raw.get("risk_level") or _risk(ai_score)).upper(),
        "indicators": list(ai_raw.get("indicators") or [])[:8],
        "reasoning": str(ai_raw.get("reasoning") or fallback["ai_explanation"]),
        "method": "hybrid_llm"
    }
    scam = {
        "score": scam_score,
        "risk_level": str(scam_raw.get("risk_level") or _risk(scam_score)).upper(),
        "indicators": list(scam_raw.get("indicators") or [])[:8],
        "reasoning": str(scam_raw.get("reasoning") or fallback["scam_explanation"]),
        "method": "hybrid_llm"
    }
    return _compact_result(ai, scam, llm_used=True, llm_status="used")


def analyze_with_llm(text):
    global _llm_unavailable, _last_llm_error
    fallback = rule_based_analysis(text, "rules")
    chronology = _extract_chronology(text)
    today = date.today()

    if not _llm_enabled():
        status = "disabled" if os.environ.get("SCAM_FILTER_USE_LLM", "auto").lower() in ("0", "false", "no", "off", "disabled") else "no_api_key"
        fallback["llm_status"] = status
        fallback["chronology_summary"] = chronology["summary"]
        return _sanitize_chronology_claims(fallback, chronology)
    if _llm_unavailable:
        fallback["llm_status"] = "unavailable_after_error"
        fallback["chronology_summary"] = chronology["summary"]
        return _sanitize_chronology_claims(fallback, chronology)

    cache_key = hashlib.sha256(
        f"{PROMPT_VERSION}\n{today.isoformat()}\n{chronology['summary']}\n{text[:MAX_LLM_CHARS]}".encode("utf-8", errors="ignore")
    ).hexdigest()
    if cache_key in _llm_cache:
        return _llm_cache[cache_key]

    prompt = f"""
You are a strict resume fraud and AI-writing detector.
Today's date is {today.isoformat()}.

Return ONLY valid JSON:
{{
  "ai_generated_percentage": 0,
  "ai_risk": "LOW",
  "ai_explanation": "short evidence-based reason",
  "scam_percentage": 0,
  "scam_risk": "LOW",
  "scam_explanation": "short evidence-based reason"
}}

Rules:
- Scores are 0-100.
- Risk must be LOW, MEDIUM, or HIGH.
- Scam means fake/fabricated applicant signals, suspicious claims, or unverifiable/fraud patterns.
- AI-generated means likely AI-written or heavily AI-polished text.
- Do not mark a normal student resume as scam just because it is short.
- Missing phone/email is a verification issue, not automatic scam.
- Do not claim a timeline is suspicious unless the text shows an actual contradiction.
- A student can validly have internships, part-time work, or projects during an ongoing degree.
- An education start date followed by an internship in a later semester or summer is normal and not suspicious.
- Only flag chronology if dates are impossible on their face, clearly overlap in a contradictory way, or are in the future relative to the resume context.
- Use this parsed chronology guidance when evaluating dates: {chronology["summary"]}
- Mention concrete evidence from the text.

Resume:
\"\"\"
{text[:MAX_LLM_CHARS]}
\"\"\"
"""
    try:
        response = _get_client().models.generate_content(
            model=LLM_MODEL,
            contents=prompt
        )
        parsed = json.loads(_strip_json_fence(response.text))
        result = _normalize_llm_result(parsed, fallback)
        result["chronology_summary"] = chronology["summary"]
        result = _sanitize_chronology_claims(result, chronology)
        _llm_cache[cache_key] = result
        return result
    except Exception as e:
        _last_llm_error = str(e)
        if _is_llm_error(e):
            _llm_unavailable = True
        fallback["llm_status"] = "fallback_after_llm_error"
        fallback["chronology_summary"] = chronology["summary"]
        fallback["ai_explanation"] = (
            "LLM analysis was unavailable, so this is a rule-based estimate from generic phrasing, "
            "repetition, vocabulary variety, and template markers."
        )
        fallback["scam_explanation"] = (
            "LLM analysis was unavailable, so this is a rule-based check using contact info, "
            "resume sections, length, links, and suspicious hiring/fraud phrases."
        )
        if os.environ.get("SCAM_FILTER_DEBUG", "").lower() in ("1", "true", "yes"):
            fallback["llm_error"] = _last_llm_error
        return _sanitize_chronology_claims(fallback, chronology)


def run_pipeline(file_path):
    if not os.path.exists(file_path):
        return {"error": "File not found"}

    text = extract_text(file_path)
    if not text or text.startswith("ERROR"):
        return {"error": text}

    return analyze_with_llm(text)


if __name__ == "__main__":
    file_path = r"D:\job-align\ml\data\raw\ai_resume_test.pdf"
    result = run_pipeline(file_path)
    print(json.dumps(result, indent=2))
